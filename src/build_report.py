"""Gera relatorio_vendas.xlsx a partir de src/sales.db.

Le o banco com pandas, calcula metricas de negocio e escreve um Excel com
duas abas: "Resumo" (KPIs + top produtos + grafico) e "Detalhe" (um pedido
por linha). Rode src/generate_db.py antes deste script.

Uso:
    python src/build_report.py
"""
import os
import sqlite3

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DB_PATH = os.path.join(os.path.dirname(__file__), "sales.db")
REPORT_PATH = os.path.join(os.path.dirname(__file__), "relatorio_vendas.xlsx")

DETAIL_QUERY = """
SELECT
    o.order_id,
    o.order_date,
    c.name AS customer_name,
    c.city AS customer_city,
    p.name AS product_name,
    p.category AS product_category,
    o.quantity,
    o.total
FROM orders o
JOIN customers c ON c.customer_id = o.customer_id
JOIN products p ON p.product_id = o.product_id
ORDER BY o.order_date
"""


def load_detail(db_path=DB_PATH) -> pd.DataFrame:
    conn = sqlite3.connect(db_path)
    try:
        return pd.read_sql_query(DETAIL_QUERY, conn)
    finally:
        conn.close()


def top_products(detail: pd.DataFrame, n=10) -> pd.DataFrame:
    return (
        detail.groupby("product_name")["total"]
        .sum()
        .sort_values(ascending=False)
        .head(n)
        .reset_index()
        .rename(columns={"product_name": "Produto", "total": "Receita"})
    )


def autosize_columns(ws, df, start_col=1):
    for i, col in enumerate(df.columns, start=start_col):
        width = max(len(str(col)), df[col].astype(str).map(len).max() if len(df) else 0) + 2
        ws.column_dimensions[get_column_letter(i)].width = width


def write_header(ws, row, cols, start_col=1):
    for i, col in enumerate(cols, start=start_col):
        cell = ws.cell(row=row, column=i, value=col)
        cell.font = Font(bold=True)


def build_report(db_path=DB_PATH, report_path=REPORT_PATH) -> str:
    detail = load_detail(db_path)
    if detail.empty:
        raise ValueError("Nenhum pedido encontrado em sales.db - rode generate_db.py primeiro.")

    total_revenue = detail["total"].sum()
    total_orders = detail["order_id"].nunique()
    avg_ticket = total_revenue / total_orders
    top = top_products(detail)

    wb = Workbook()

    # --- Aba Resumo ---
    ws_summary = wb.active
    ws_summary.title = "Resumo"
    ws_summary["A1"] = "Resumo de Vendas"
    ws_summary["A1"].font = Font(bold=True, size=14)

    kpis = [
        ("Receita total", round(total_revenue, 2)),
        ("Total de pedidos", total_orders),
        ("Ticket medio", round(avg_ticket, 2)),
    ]
    for i, (label, value) in enumerate(kpis, start=3):
        ws_summary.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=i, column=2, value=value)

    top_start_row = 3 + len(kpis) + 2
    ws_summary.cell(row=top_start_row - 1, column=1, value="Top produtos por receita").font = Font(bold=True, size=12)
    write_header(ws_summary, top_start_row, top.columns)
    for r, (_, row) in enumerate(top.iterrows(), start=top_start_row + 1):
        ws_summary.cell(row=r, column=1, value=row["Produto"])
        ws_summary.cell(row=r, column=2, value=round(row["Receita"], 2))
    autosize_columns(ws_summary, top)

    chart = BarChart()
    chart.title = "Top produtos por receita"
    chart.y_axis.title = "Receita"
    data = Reference(ws_summary, min_col=2, min_row=top_start_row, max_row=top_start_row + len(top))
    categories = Reference(ws_summary, min_col=1, min_row=top_start_row + 1, max_row=top_start_row + len(top))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws_summary.add_chart(chart, f"D{top_start_row - 1}")

    # --- Aba Detalhe ---
    ws_detail = wb.create_sheet("Detalhe")
    write_header(ws_detail, 1, detail.columns)
    for r, (_, row) in enumerate(detail.iterrows(), start=2):
        for c, value in enumerate(row, start=1):
            ws_detail.cell(row=r, column=c, value=value)
    autosize_columns(ws_detail, detail)

    wb.save(report_path)
    return report_path


if __name__ == "__main__":
    path = build_report()
    print(f"Relatorio gerado em {path}")
