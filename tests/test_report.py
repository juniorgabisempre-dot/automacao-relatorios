"""Teste de sanidade: gera o DB e o relatorio em um diretorio temporario e
confere que o Excel tem as abas esperadas e linhas de dados."""
import os
import sqlite3
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from openpyxl import load_workbook  # noqa: E402

import generate_db  # noqa: E402
from build_report import build_report  # noqa: E402


class TestBuildReport(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.db_path = os.path.join(self.tmpdir.name, "sales.db")
        self.report_path = os.path.join(self.tmpdir.name, "relatorio_vendas.xlsx")

        conn = sqlite3.connect(self.db_path)
        generate_db.build_schema(conn)
        import random

        from faker import Faker

        random.seed(1)
        fake = Faker("pt_BR")
        Faker.seed(1)
        generate_db.generate(conn, fake)
        conn.commit()
        conn.close()

    def tearDown(self):
        self.tmpdir.cleanup()

    def test_report_has_expected_sheets_and_rows(self):
        build_report(db_path=self.db_path, report_path=self.report_path)
        self.assertTrue(os.path.exists(self.report_path))

        wb = load_workbook(self.report_path)
        self.assertEqual(set(wb.sheetnames), {"Resumo", "Detalhe"})

        detail_rows = wb["Detalhe"].max_row
        self.assertGreater(detail_rows, 1)  # header + pelo menos 1 pedido


if __name__ == "__main__":
    unittest.main()
